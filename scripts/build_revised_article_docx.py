from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/wangjiangyu/Documents/Teaching Law")
SOURCE = ROOT / "docs/drafts/teaching-iel-ai-geopolitics-revised.md"
OUTPUT = Path(
    "/Users/wangjiangyu/Library/CloudStorage/OneDrive-CityUniversityofHongKong/"
    "0-MyWorks/Teaching IEL in the age of AI and Geopolitics - revised.docx"
)


FOOTNOTE_DEF_RE = re.compile(r"^\[\^([^\]]+)\]:\s*(.*)$")
FOOTNOTE_REF_RE = re.compile(r"\[\^([^\]]+)\]")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "C9CDD3")


def parse_markdown(path: Path):
    text = path.read_text(encoding="utf-8")
    footnotes: dict[str, str] = {}
    main_lines: list[str] = []
    in_notes = False
    for line in text.splitlines():
        if line.strip() == "## References and Source Notes":
            in_notes = True
            continue
        if in_notes:
            match = FOOTNOTE_DEF_RE.match(line)
            if match:
                footnotes[match.group(1)] = match.group(2).strip()
            elif footnotes and line.strip():
                last_key = next(reversed(footnotes))
                footnotes[last_key] += " " + line.strip()
            continue
        main_lines.append(line.rstrip())
    return main_lines, footnotes


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(8)

    for name, size, color, before, after in [
        ("Heading 1", 16, "1F4E79", 14, 8),
        ("Heading 2", 14, "1F4E79", 12, 6),
        ("Heading 3", 12, "365F91", 10, 4),
    ]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_runs_with_citations(paragraph, text: str, footnotes: dict[str, str], note_numbers: dict[str, int], note_order: list[str]) -> None:
    pos = 0
    for match in FOOTNOTE_REF_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        key = match.group(1)
        if key not in note_numbers:
            note_numbers[key] = len(note_order) + 1
            note_order.append(key)
        run = paragraph.add_run(str(note_numbers[key]))
        run.font.superscript = True
        run.font.size = Pt(9)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def add_paragraph(doc: Document, text: str, footnotes, note_numbers, note_order, style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_runs_with_citations(paragraph, text, footnotes, note_numbers, note_order)
    return paragraph


def add_title_block(doc: Document, title: str, subtitle: str) -> None:
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(3)
    title_run = title_p.add_run(title)
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor.from_string("0B2545")

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_p.paragraph_format.space_after = Pt(16)
    subtitle_run = subtitle_p.add_run(subtitle)
    subtitle_run.italic = True
    subtitle_run.font.name = "Times New Roman"
    subtitle_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    subtitle_run.font.size = Pt(13)
    subtitle_run.font.color.rgb = RGBColor.from_string("365F91")


def split_table_row(line: str) -> list[str]:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return cells


def add_markdown_table(doc: Document, rows: list[str], footnotes, note_numbers, note_order) -> None:
    data = [split_table_row(row) for row in rows if not re.match(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$", row)]
    if not data:
        return
    col_count = max(len(row) for row in data)
    table = doc.add_table(rows=len(data), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx]
        if r_idx == 0:
            set_repeat_table_header(row)
        for c_idx in range(col_count):
            cell = row.cells[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_shading(cell, "E8EEF5")
            text = row_data[c_idx] if c_idx < len(row_data) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_runs_with_citations(paragraph, text, footnotes, note_numbers, note_order)
            for run in paragraph.runs:
                run.font.size = Pt(10)
                if r_idx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string("0B2545")
    doc.add_paragraph()


def add_notes(doc: Document, footnotes: dict[str, str], note_order: list[str]) -> None:
    doc.add_page_break()
    doc.add_heading("References and Source Notes", level=1)
    for i, key in enumerate(note_order, start=1):
        text = footnotes.get(key, f"Missing source note: {key}")
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(6)
        number_run = paragraph.add_run(f"{i}. ")
        number_run.bold = True
        number_run.font.name = "Times New Roman"
        number_run.font.size = Pt(10)
        body_run = paragraph.add_run(text)
        body_run.font.name = "Times New Roman"
        body_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        body_run.font.size = Pt(10)


def build() -> None:
    lines, footnotes = parse_markdown(SOURCE)
    doc = Document()
    configure_document(doc)

    note_numbers: dict[str, int] = {}
    note_order: list[str] = []

    title = None
    subtitle = None
    idx = 0
    while idx < len(lines):
        line = lines[idx].strip()
        if line.startswith("# "):
            title = line[2:].strip()
        elif line.startswith("## "):
            subtitle = line[3:].strip()
            idx += 1
            break
        idx += 1
    add_title_block(doc, title or "Untitled", subtitle or "")

    table_buffer: list[str] = []
    paragraph_buffer: list[str] = []

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            text = " ".join(part.strip() for part in paragraph_buffer).strip()
            if text:
                if re.match(r"^\d+\.\s+", text):
                    add_paragraph(doc, re.sub(r"^\d+\.\s+", "", text), footnotes, note_numbers, note_order, style="List Number")
                else:
                    add_paragraph(doc, text, footnotes, note_numbers, note_order)
            paragraph_buffer = []

    def flush_table():
        nonlocal table_buffer
        if table_buffer:
            add_markdown_table(doc, table_buffer, footnotes, note_numbers, note_order)
            table_buffer = []

    for raw in lines[idx:]:
        line = raw.rstrip()
        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            flush_table()
            continue
        if stripped.startswith("|"):
            flush_paragraph()
            table_buffer.append(stripped)
            continue
        flush_table()
        if stripped.startswith("### "):
            flush_paragraph()
            add_paragraph(doc, stripped[4:], footnotes, note_numbers, note_order, style="Heading 2")
        elif stripped.startswith("## "):
            flush_paragraph()
            add_paragraph(doc, stripped[3:], footnotes, note_numbers, note_order, style="Heading 1")
        elif stripped.startswith("# "):
            flush_paragraph()
        else:
            paragraph_buffer.append(stripped)

    flush_paragraph()
    flush_table()
    add_notes(doc, footnotes, note_order)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
