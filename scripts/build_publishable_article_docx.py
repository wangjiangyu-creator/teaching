from __future__ import annotations

import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree


ROOT = Path("/Users/wangjiangyu/Documents/Teaching Law")
SOURCE = ROOT / "docs/drafts/teaching-iel-ai-geopolitics-publishable.md"
WORKSPACE_OUTPUT = ROOT / "deliverables/Teaching IEL after AI and Geoeconomic Fragmentation - JLE submission.docx"
ONEDRIVE_OUTPUT = Path(
    "/Users/wangjiangyu/Library/CloudStorage/OneDrive-CityUniversityofHongKong/"
    "0-MyWorks/Teaching IEL after AI and Geoeconomic Fragmentation - JLE submission.docx"
)

FOOTNOTE_DEF_RE = re.compile(r"^\[\^([^\]]+)\]:\s*(.*)$")
FOOTNOTE_REF_RE = re.compile(r"\[\^([^\]]+)\]")
TABLE_SEPARATOR_RE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$")

FONT = "Times New Roman"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKGREL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
NS = {"w": W_NS, "r": R_NS, "rel": PKGREL_NS, "ct": CT_NS}
REL_TYPE_FOOTNOTES = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"
CT_FOOTNOTES = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"


def set_font(run, size: float | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, before: float = 0, after: float = 0, line_spacing: float = 2.0) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line_spacing


def parse_markdown(path: Path) -> tuple[list[str], dict[str, str]]:
    text = path.read_text(encoding="utf-8")
    footnotes: dict[str, str] = {}
    main_lines: list[str] = []
    in_notes = False
    for line in text.splitlines():
        if line.strip() == "## References and Source Notes":
            in_notes = True
            continue
        if in_notes:
            match = FOOTNOTE_DEF_RE.match(line)
            if match:
                footnotes[match.group(1)] = match.group(2).strip()
            elif footnotes and line.strip():
                last_key = next(reversed(footnotes))
                footnotes[last_key] += " " + line.strip()
            continue
        main_lines.append(line.rstrip())
    return main_lines, footnotes


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 2.0

    style_tokens = {
        "Heading 1": (12, "000000", 12, 6),
        "Heading 2": (12, "000000", 12, 6),
        "Heading 3": (12, "000000", 12, 6),
    }
    for name, (size, color, before, after) in style_tokens.items():
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 2.0
        style.paragraph_format.keep_with_next = True

    for list_style_name in ("List Number", "List Bullet"):
        style = styles[list_style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(12)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 2.0
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)

    for note_style_name in ("Footnote Text", "Footnote Reference"):
        try:
            style = styles[note_style_name]
        except KeyError:
            continue
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(10)
        if note_style_name == "Footnote Text":
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.space_after = Pt(0)
            style.paragraph_format.line_spacing = 1.0


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_separate)
    run._r.append(text)
    run._r.append(fld_char_end)
    set_font(run, 12)


def add_footer(doc: Document) -> None:
    footer = doc.sections[0].footer
    paragraph = footer.paragraphs[0]
    add_page_number(paragraph)


def add_runs_with_citations(paragraph, text: str, footnotes: dict[str, str], note_order: list[str], size: float = 12) -> None:
    pos = 0
    for match in FOOTNOTE_REF_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_font(run, size)
        key = match.group(1)
        if key not in footnotes:
            raise KeyError(f"Missing footnote definition for {key}")
        note_order.append(key)
        run = paragraph.add_run(f"[[FN{len(note_order):04d}]]")
        set_font(run, size)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_font(run, size)


def add_paragraph(
    doc: Document,
    text: str,
    footnotes: dict[str, str],
    note_order: list[str],
    style: str | None = None,
    size: float = 12,
    justify: bool = True,
):
    paragraph = doc.add_paragraph(style=style)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_runs_with_citations(paragraph, text, footnotes, note_order, size=size)
    if not style:
        set_paragraph_spacing(paragraph)
    return paragraph


def add_title_block(doc: Document, title: str, subtitle: str) -> None:
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(0)
    title_p.paragraph_format.line_spacing = 2.0
    title_run = title_p.add_run(title)
    set_font(title_run, 12, bold=True)
    title_run.font.color.rgb = RGBColor.from_string("000000")

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_p.paragraph_format.space_after = Pt(0)
    subtitle_p.paragraph_format.line_spacing = 2.0
    subtitle_run = subtitle_p.add_run(subtitle)
    set_font(subtitle_run, 12, italic=True)
    subtitle_run.font.color.rgb = RGBColor.from_string("000000")


def split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "C9CDD3")


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)


def widths_for_table(col_count: int) -> list[int]:
    if col_count == 4:
        return [1980, 2460, 2040, 2880]
    if col_count == 3:
        return [2220, 3570, 3570]
    if col_count == 2:
        return [2400, 6960]
    base = CONTENT_WIDTH_DXA // col_count
    widths = [base] * col_count
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def add_markdown_table(doc: Document, rows: list[str], footnotes: dict[str, str], note_order: list[str]) -> None:
    data = [split_table_row(row) for row in rows if not TABLE_SEPARATOR_RE.match(row)]
    if not data:
        return
    doc.add_page_break()
    col_count = max(len(row) for row in data)
    widths = widths_for_table(col_count)
    table = doc.add_table(rows=len(data), cols=col_count)
    set_table_geometry(table, widths)
    set_table_borders(table)
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx]
        if r_idx == 0:
            set_repeat_table_header(row)
        for c_idx in range(col_count):
            cell = row.cells[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cell, widths[c_idx])
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_shading(cell, "F2F4F7")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            text = row_data[c_idx] if c_idx < len(row_data) else ""
            add_runs_with_citations(paragraph, text, footnotes, note_order, size=10)
            for run in paragraph.runs:
                if r_idx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string("000000")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def xml_bytes(root: etree._Element) -> bytes:
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")


def read_xml(z: zipfile.ZipFile, name: str) -> etree._Element:
    return etree.fromstring(z.read(name))


def next_rid(rels_root: etree._Element) -> str:
    max_n = 0
    for rel in rels_root.findall(f"{{{PKGREL_NS}}}Relationship"):
        rid = rel.get("Id") or ""
        match = re.match(r"rId(\d+)$", rid)
        if match:
            max_n = max(max_n, int(match.group(1)))
    return f"rId{max_n + 1}"


def ensure_footnote_relationship(rels_root: etree._Element) -> bool:
    for rel in rels_root.findall(f"{{{PKGREL_NS}}}Relationship"):
        if rel.get("Type") == REL_TYPE_FOOTNOTES:
            return False
    rel = etree.SubElement(rels_root, f"{{{PKGREL_NS}}}Relationship")
    rel.set("Id", next_rid(rels_root))
    rel.set("Type", REL_TYPE_FOOTNOTES)
    rel.set("Target", "footnotes.xml")
    return True


def ensure_content_type_override(ct_root: etree._Element) -> bool:
    part_name = "/word/footnotes.xml"
    for override in ct_root.findall(f"{{{CT_NS}}}Override"):
        if override.get("PartName") == part_name:
            if override.get("ContentType") != CT_FOOTNOTES:
                override.set("ContentType", CT_FOOTNOTES)
                return True
            return False
    override = etree.SubElement(ct_root, f"{{{CT_NS}}}Override")
    override.set("PartName", part_name)
    override.set("ContentType", CT_FOOTNOTES)
    return True


def make_footnotes_part(note_order: list[str], footnotes: dict[str, str]) -> etree._Element:
    root = etree.Element(f"{{{W_NS}}}footnotes", nsmap={"w": W_NS, "r": R_NS})
    for note_id, sep_tag in (("-1", "separator"), ("0", "continuationSeparator")):
        note = etree.SubElement(root, f"{{{W_NS}}}footnote")
        note.set(f"{{{W_NS}}}id", note_id)
        paragraph = etree.SubElement(note, f"{{{W_NS}}}p")
        run = etree.SubElement(paragraph, f"{{{W_NS}}}r")
        etree.SubElement(run, f"{{{W_NS}}}{sep_tag}")

    for idx, key in enumerate(note_order, start=1):
        note = etree.SubElement(root, f"{{{W_NS}}}footnote")
        note.set(f"{{{W_NS}}}id", str(idx))
        paragraph = etree.SubElement(note, f"{{{W_NS}}}p")
        p_pr = etree.SubElement(paragraph, f"{{{W_NS}}}pPr")
        p_style = etree.SubElement(p_pr, f"{{{W_NS}}}pStyle")
        p_style.set(f"{{{W_NS}}}val", "FootnoteText")

        ref_run = etree.SubElement(paragraph, f"{{{W_NS}}}r")
        r_pr = etree.SubElement(ref_run, f"{{{W_NS}}}rPr")
        r_style = etree.SubElement(r_pr, f"{{{W_NS}}}rStyle")
        r_style.set(f"{{{W_NS}}}val", "FootnoteReference")
        etree.SubElement(ref_run, f"{{{W_NS}}}footnoteRef")

        space_run = etree.SubElement(paragraph, f"{{{W_NS}}}r")
        space_text = etree.SubElement(space_run, f"{{{W_NS}}}t")
        space_text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        space_text.text = " "

        text_run = etree.SubElement(paragraph, f"{{{W_NS}}}r")
        text = etree.SubElement(text_run, f"{{{W_NS}}}t")
        text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        text.text = footnotes[key]
    return root


def replace_footnote_markers(document_root: etree._Element, note_count: int) -> None:
    for idx in range(1, note_count + 1):
        marker = f"[[FN{idx:04d}]]"
        replaced = False
        for text_node in document_root.xpath(".//w:t", namespaces=NS):
            if text_node.text != marker:
                continue
            run = text_node.getparent()
            while run is not None and run.tag != f"{{{W_NS}}}r":
                run = run.getparent()
            if run is None:
                continue
            for child in list(run):
                run.remove(child)
            r_pr = etree.SubElement(run, f"{{{W_NS}}}rPr")
            r_style = etree.SubElement(r_pr, f"{{{W_NS}}}rStyle")
            r_style.set(f"{{{W_NS}}}val", "FootnoteReference")
            ref = etree.SubElement(run, f"{{{W_NS}}}footnoteReference")
            ref.set(f"{{{W_NS}}}id", str(idx))
            replaced = True
            break
        if not replaced:
            raise RuntimeError(f"Could not replace footnote marker {marker}")


def patch_true_footnotes(in_docx: Path, out_docx: Path, note_order: list[str], footnotes: dict[str, str]) -> None:
    with zipfile.ZipFile(in_docx, "r") as zin:
        document_root = read_xml(zin, "word/document.xml")
        replace_footnote_markers(document_root, len(note_order))

        rels_root = read_xml(zin, "word/_rels/document.xml.rels")
        rels_changed = ensure_footnote_relationship(rels_root)

        ct_root = read_xml(zin, "[Content_Types].xml")
        ct_changed = ensure_content_type_override(ct_root)

        footnotes_root = make_footnotes_part(note_order, footnotes)
        overrides: dict[str, bytes] = {
            "word/document.xml": xml_bytes(document_root),
            "word/footnotes.xml": xml_bytes(footnotes_root),
        }
        if rels_changed:
            overrides["word/_rels/document.xml.rels"] = xml_bytes(rels_root)
        if ct_changed:
            overrides["[Content_Types].xml"] = xml_bytes(ct_root)

        original_names = {info.filename for info in zin.infolist()}
        with zipfile.ZipFile(out_docx, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                if info.filename in overrides:
                    zout.writestr(info.filename, overrides[info.filename])
                elif info.filename != "word/footnotes.xml":
                    zout.writestr(info.filename, zin.read(info.filename))
            for name, data in overrides.items():
                if name not in original_names:
                    zout.writestr(name, data)


def add_notes(doc: Document, footnotes: dict[str, str], note_order: list[str]) -> None:
    doc.add_page_break()
    heading = doc.add_paragraph(style="Heading 1")
    heading.add_run("References and Source Notes")
    for run in heading.runs:
        set_font(run, 16, bold=True)
        run.font.color.rgb = RGBColor.from_string("2E74B5")
    for i, key in enumerate(note_order, start=1):
        text = footnotes.get(key, f"Missing source note: {key}")
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.15
        number_run = paragraph.add_run(f"{i}. ")
        set_font(number_run, 9.5, bold=True)
        body_run = paragraph.add_run(text)
        set_font(body_run, 9.5)


def build() -> None:
    lines, footnotes = parse_markdown(SOURCE)
    doc = Document()
    configure_document(doc)
    add_footer(doc)

    note_order: list[str] = []

    title = ""
    subtitle = ""
    idx = 0
    while idx < len(lines):
        stripped = lines[idx].strip()
        if stripped.startswith("# "):
            title = stripped[2:].strip()
        elif stripped.startswith("## "):
            subtitle = stripped[3:].strip()
            idx += 1
            break
        idx += 1
    add_title_block(doc, title or "Untitled", subtitle)

    table_buffer: list[str] = []
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            text = " ".join(part.strip() for part in paragraph_buffer).strip()
            if text:
                if re.match(r"^\d+\.\s+", text):
                    add_paragraph(doc, re.sub(r"^\d+\.\s+", "", text), footnotes, note_order, style="List Number")
                elif text.startswith("- "):
                    add_paragraph(doc, text[2:], footnotes, note_order, style="List Bullet")
                else:
                    add_paragraph(doc, text, footnotes, note_order)
            paragraph_buffer = []

    def flush_table() -> None:
        nonlocal table_buffer
        if table_buffer:
            add_markdown_table(doc, table_buffer, footnotes, note_order)
            table_buffer = []

    for raw in lines[idx:]:
        line = raw.rstrip()
        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            flush_table()
            continue
        if stripped.startswith("|"):
            flush_paragraph()
            table_buffer.append(stripped)
            continue
        flush_table()
        if stripped.startswith("### "):
            flush_paragraph()
            add_paragraph(doc, stripped[4:], footnotes, note_order, style="Heading 2", justify=False)
        elif stripped.startswith("## "):
            flush_paragraph()
            add_paragraph(doc, stripped[3:], footnotes, note_order, style="Heading 1", justify=False)
        elif stripped.startswith("# "):
            flush_paragraph()
        else:
            paragraph_buffer.append(stripped)

    flush_paragraph()
    flush_table()
    temp_output = WORKSPACE_OUTPUT.with_name(WORKSPACE_OUTPUT.stem + " - unpatched.docx")
    temp_output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(temp_output)

    WORKSPACE_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    patch_true_footnotes(temp_output, WORKSPACE_OUTPUT, note_order, footnotes)
    temp_output.unlink(missing_ok=True)
    print(WORKSPACE_OUTPUT)

    ONEDRIVE_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(WORKSPACE_OUTPUT, ONEDRIVE_OUTPUT)
    print(ONEDRIVE_OUTPUT)


if __name__ == "__main__":
    build()
